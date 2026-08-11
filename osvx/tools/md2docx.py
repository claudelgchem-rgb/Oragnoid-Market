#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
OSVX — Markdown -> Word(.docx) 변환기 (Agent W 지원 도구)

report/OSVX_report.md 를 실제 Word 표·제목 스타일·목차 필드를 가진 .docx 로 변환한다.
GFM 파이프 표 -> 진짜 Word 표, 헤딩 -> Word 헤딩(탐색창/목차 연동), 코드블록 -> 음영 고정폭.

usage: python3 tools/md2docx.py [입력.md] [출력.docx]
기본값: report/OSVX_report.md -> report/OSVX_report.docx
"""
import os
import re
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KO_FONT = "Malgun Gothic"
MONO = "Consolas"
ACCENT = RGBColor(0x1F, 0x3B, 0x63)
MUTE = RGBColor(0x5D, 0x66, 0x73)


# ------------------------------------------------------------------ helpers
def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def style_font(run, size=None, bold=None, italic=None, color=None, mono=False):
    f = run.font
    name = MONO if mono else KO_FONT
    f.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


INLINE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|`[^`]+`|\*[^*\n]+\*|~~.+?~~)")


def add_inline(par, text, size=10.5, base_bold=False, color=None):
    """굵게/기울임/코드 인라인 마크업을 Word run 으로."""
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)  # 링크 평문화
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("***") and tok.endswith("***"):
            r = par.add_run(tok[3:-3]); style_font(r, size, True, True, color)
        elif tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2]); style_font(r, size, True, None, color)
        elif tok.startswith("~~") and tok.endswith("~~"):
            r = par.add_run(tok[2:-2]); style_font(r, size, base_bold, None, MUTE)
            r.font.strike = True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 1:
            r = par.add_run(tok[1:-1]); style_font(r, size - 0.5, None, None, RGBColor(0xA3, 0x1D, 0x1D), mono=True)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 1:
            r = par.add_run(tok[1:-1]); style_font(r, size, base_bold, True, color)
        else:
            r = par.add_run(tok); style_font(r, size, base_bold, None, color)


def is_table_line(l):
    return l.strip().startswith("|") and l.count("|") >= 2


def is_sep_line(l):
    return bool(re.match(r"^\s*\|[\s\-:|]+\|\s*$", l))


def split_row(l):
    s = l.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


# ------------------------------------------------------------------ builders
def base_styles(doc):
    st = doc.styles["Normal"]
    st.font.name = KO_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), KO_FONT)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.line_spacing = 1.32
    for i, (sz, col) in enumerate([(19, ACCENT), (14.5, ACCENT), (12, ACCENT), (11, MUTE)], start=1):
        try:
            h = doc.styles["Heading %d" % i]
        except KeyError:
            continue
        h.font.name = KO_FONT
        h.font.size = Pt(sz)
        h.font.bold = True
        h.font.color.rgb = col
        h.element.rPr.rFonts.set(qn("w:eastAsia"), KO_FONT)
        h.paragraph_format.space_before = Pt(16 if i <= 2 else 10)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True


def add_toc(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "목차를 갱신하려면 이 영역을 선택하고 F9 를 누르십시오."
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for e in (fld, instr, sep, t, end):
        r._r.append(e)


def add_footer_pagenum(doc):
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        for typ, txt in (("begin", None), (None, "PAGE"), ("separate", None), ("text", "1"), ("end", None)):
            if typ == "text":
                e = OxmlElement("w:t"); e.text = txt
            elif txt == "PAGE":
                e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = " PAGE "
            else:
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), typ)
            r._r.append(e)
        style_font(r, 9, color=MUTE)


def make_table(doc, rows):
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            c = t.cell(ri, ci)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.12
            add_inline(p, cell, size=8.8, base_bold=(ri == 0))
            if ri == 0:
                set_cell_bg(c, "1F3B63")
                for r_ in p.runs:
                    r_.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    r_.font.bold = True
            elif ri % 2 == 0:
                set_cell_bg(c, "F2F5F9")
    set_repeat_header(t.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return t


def convert(src, dst):
    with open(src, "r", encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    base_styles(doc)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    i = 0
    n = len(lines)
    in_code = False
    code_buf = []
    made_toc = False

    while i < n:
        line = lines[i]
        s = line.strip()

        # 코드블록
        if s.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                r = p.add_run("\n".join(code_buf))
                style_font(r, 8.8, mono=True)
                set_par_bg(p, "F4F6F8")
                code_buf, in_code = [], False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # 페이지 나눔
        if s in ("<!--pagebreak-->", "\\pagebreak", "---PAGEBREAK---"):
            doc.add_page_break()
            i += 1
            continue

        # 목차 자리표시자
        if s in ("<!--toc-->", "[TOC]"):
            add_toc(doc)
            made_toc = True
            i += 1
            continue

        # 표
        if is_table_line(line):
            block = []
            while i < n and is_table_line(lines[i]):
                if not is_sep_line(lines[i]):
                    block.append(split_row(lines[i]))
                i += 1
            if block:
                make_table(doc, block)
            continue

        # 수평선
        if re.match(r"^\s*(\*\s*){3,}$|^\s*(-\s*){3,}$|^\s*(_\s*){3,}$", line):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bt = OxmlElement("w:bottom")
            bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "6")
            bt.set(qn("w:space"), "1"); bt.set(qn("w:color"), "C6CDD6")
            pbdr.append(bt); pPr.append(pbdr)
            i += 1
            continue

        # 헤딩
        m = re.match(r"^(#{1,5})\s+(.*)$", s)
        if m:
            lvl, txt = len(m.group(1)), m.group(2).strip()
            txt = re.sub(r"\s*#+\s*$", "", txt)
            if lvl == 1 and not made_toc and doc.paragraphs and len(doc.paragraphs) > 3:
                doc.add_page_break()
            h = doc.add_heading(level=min(lvl, 4))
            h.text = ""
            add_inline(h, txt, size={1: 19, 2: 14.5, 3: 12, 4: 11}.get(lvl, 10.5), base_bold=True,
                       color=ACCENT if lvl <= 3 else MUTE)
            i += 1
            continue

        # 인용
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            add_inline(p, s.lstrip("> ").strip(), size=10, color=MUTE)
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            lf = OxmlElement("w:left")
            lf.set(qn("w:val"), "single"); lf.set(qn("w:sz"), "18")
            lf.set(qn("w:space"), "8"); lf.set(qn("w:color"), "1F3B63")
            pbdr.append(lf); pPr.append(pbdr)
            i += 1
            continue

        # 목록
        mb = re.match(r"^(\s*)([-*+])\s+(.*)$", line)
        mn = re.match(r"^(\s*)(\d+)[\.\)]\s+(.*)$", line)
        if mb or mn:
            mm = mb or mn
            indent = len(mm.group(1)) // 2
            style = "List Bullet" if mb else "List Number"
            try:
                p = doc.add_paragraph(style=style if indent == 0 else "%s %d" % (style, min(indent + 1, 3)))
            except KeyError:
                p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, mm.group(3))
            i += 1
            continue

        # 빈 줄
        if not s:
            i += 1
            continue

        # 일반 문단
        p = doc.add_paragraph()
        add_inline(p, s)
        i += 1

    add_footer_pagenum(doc)
    doc.save(dst)
    return doc


def set_par_bg(p, hexcolor):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    pPr.append(shd)


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "report", "OSVX_report.md")
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, "report", "OSVX_report.docx")
    if not os.path.exists(src):
        print("입력 없음: %s" % src, file=sys.stderr)
        return 1
    d = convert(src, dst)
    nt = len(d.tables)
    npar = len(d.paragraphs)
    print("변환 완료: %s -> %s" % (os.path.relpath(src, ROOT), os.path.relpath(dst, ROOT)))
    print("  Word 표 %d개 / 문단 %d개 / %.1f KB" % (nt, npar, os.path.getsize(dst) / 1024.0))
    return 0


if __name__ == "__main__":
    sys.exit(main())
